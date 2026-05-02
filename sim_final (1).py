import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import sys, os

# ── PALETTE ───────────────────────────────────────────────────────────────────
C = {
    "navy":         "1B2A4A",
    "navy_mid":     "2C3E6B",
    "header_sub":   "E8EDF5",
    "white":        "FFFFFF",
    "actual_bg":    "F0F0F0",
    "input_bg":     "FFFFFF",
    "baseline_bg":  "EBF5FB",
    "var_bg":       "EAFAF1",
    "delta_bg":     "FEF9E7",
    "subtotal_bg":  "F4F6F9",
    "flag_true":    "D5F5E3",
    "flag_false":   "FDFEFE",
    "s3_title":     "1B2A4A",
    "s3_metric":    "F8F9FA",
    "s3_baseline":  "EBF5FB",
    "s3_var":       "EAFAF1",
    "s3_delta":     "FEF9E7",
    "font_white":   "FFFFFF",
    "font_dark":    "1A1A2E",
    "font_navy":    "1B2A4A",
    "font_green":   "1E8449",
    "font_red":     "C0392B",
    "font_grey":    "95A5A6",
    "font_input":   "1B2A4A",
    "border":       "D5D8DC",
    "border_dark":  "AEB6BF",
}

def fx(c):
    from openpyxl.styles import PatternFill as _PF
    return _PF("solid", fgColor=c)
def bdr(color=None):
    from openpyxl.styles import Border as _B, Side as _S
    col = color or C["border"]
    return _B(left=_S(style="thin",color=col), right=_S(style="thin",color=col),
              top=_S(style="thin",color=col),  bottom=_S(style="thin",color=col))
def ft(bold=False, color=None, size=9, italic=False, name="Calibri"):
    from openpyxl.styles import Font as _F
    return _F(name=name, bold=bold, color=color or C["font_dark"],
              size=size, italic=italic)
def al(h="left", v="center", wrap=False):
    from openpyxl.styles import Alignment as _AL
    return _AL(horizontal=h, vertical=v, wrap_text=wrap)

NUM   = '#,##0.00;(#,##0.00);"-"'
NUM0  = '#,##0;(#,##0);"-"'
PCT   = '0.0%;(0.0%);"-"'
DPCT  = '+0.0%;-0.0%;"0.0%"'
DNUM  = '+#,##0;-#,##0;"0"'
DPP   = '+0.0"pp";-0.0"pp";"0pp"'

# ── FILE PICKER ───────────────────────────────────────────────────────────────
def pick_files():
    root = tk.Tk(); root.withdraw(); root.attributes("-topmost", True)
    print("\n📂  Pilih File 1 — Data Master (semua SKU)")
    f1 = filedialog.askopenfilename(title="File 1 — Data Master",
         filetypes=[("Excel/CSV","*.xlsx *.xls *.csv")])
    if not f1: sys.exit("❌  File 1 tidak dipilih.")
    print(f"   ✅  {os.path.basename(f1)}")
    print("\n📂  Pilih File 2 — Skenario Harga")
    f2 = filedialog.askopenfilename(title="File 2 — Skenario Harga",
         filetypes=[("Excel/CSV","*.xlsx *.xls *.csv")])
    if not f2: sys.exit("❌  File 2 tidak dipilih.")
    print(f"   ✅  {os.path.basename(f2)}")
    root.destroy()
    return f1, f2

# ── LOAD ──────────────────────────────────────────────────────────────────────
class SimpleDF:
    """Minimal DataFrame-like object using list of dicts"""
    def __init__(self, rows, columns=None):
        if columns:
            self._cols = list(columns)
            self._data = [r if isinstance(r,dict) else dict(zip(columns, r)) for r in rows]
        elif rows and isinstance(rows[0], dict):
            self._cols = list(rows[0].keys())
            self._data = list(rows)
        else:
            self._cols = []
            self._data = []

    @property
    def columns(self):
        return self._cols

    def __len__(self): return len(self._data)
    def __iter__(self): return iter(self._cols)
    def __contains__(self, item): return item in self._cols
    def __getitem__(self, key):
        if isinstance(key, list):
            return SimpleDF([{c: r.get(c) for c in key} for r in self._data])
        return [r.get(key) for r in self._data]

    def __setitem__(self, key, values):
        if key not in self._cols: self._cols.append(key)
        for r, v in zip(self._data, values): r[key] = v

    def itertuples(self, index=False):
        from collections import namedtuple
        T = namedtuple("Row", [c.replace(" ","_").replace("/","_") for c in self._cols])
        for r in self._data:
            yield T(*[r.get(c) for c in self._cols])

    def iloc(self, idx): return self._data[idx]

    @property
    def shape(self): return (len(self._data), len(self._cols))

    def merge(self, other, on, how="left"):
        other_map = {r.get(on): r for r in other._data}
        result = []
        for r in self._data:
            key = r.get(on)
            merged = dict(r)
            if key in other_map:
                for k,v in other_map[key].items():
                    if k != on: merged[k] = v
            elif how == "left":
                for c in other._cols:
                    if c != on and c not in merged: merged[c] = None
            result.append(merged)
        all_cols = list(dict.fromkeys(list(self._cols) + [c for c in other._cols if c!=on]))
        return SimpleDF(result, all_cols)

    def dropna(self, subset=None):
        if subset is None: return self
        return SimpleDF([r for r in self._data
                         if all(r.get(c) is not None and str(r.get(c)).strip()!="" for c in subset)],
                        self._cols)

    def copy(self): return SimpleDF(list(self._data), list(self._cols))

    def fillna(self, val, col=None):
        result = []
        for r in self._data:
            if col:
                nr = dict(r); nr[col] = val if (r.get(col) is None or r.get(col)=="") else r[col]
                result.append(nr)
            else:
                result.append({k: (val if v is None else v) for k,v in r.items()})
        return SimpleDF(result, self._cols)

    def duplicated(self, keep=False, cols=None):
        # returns list of booleans
        col = cols[0] if cols else self._cols[0]
        from collections import Counter
        counts = Counter(r.get(col) for r in self._data)
        if keep == False:
            return [counts[r.get(col)] > 1 for r in self._data]
        seen = set()
        result = []
        for r in self._data:
            v = r.get(col)
            result.append(v in seen)
            seen.add(v)
        return result

    def any_true(self, bools): return any(bools)

    def filter_by(self, col, val):
        return SimpleDF([r for r in self._data if r.get(col)==val], self._cols)

    def unique(self, col):
        seen = []; result = []
        for r in self._data:
            v = r.get(col)
            if v not in seen: seen.append(v); result.append(v)
        return result

    def isin(self, col, vals):
        return [r.get(col) in vals for r in self._data]

    def where_not_isin(self, col, vals):
        return SimpleDF([r for r in self._data if r.get(col) not in vals], self._cols)

    def where_isin(self, col, vals):
        return SimpleDF([r for r in self._data if r.get(col) in vals], self._cols)

    def sum_col(self, col):
        return sum(r.get(col,0) or 0 for r in self._data)

    def isna_any(self, col):
        return any(r.get(col) is None or (isinstance(r.get(col),float) and r.get(col)!=r.get(col))
                   for r in self._data)


def _try_num(v):
    if v is None: return None
    try: return int(v) if str(v)==str(int(float(v))) else float(v)
    except: return v

def load_df(path):
    import csv as _csv
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            reader = _csv.DictReader(fh)
            raw = list(reader)
        if not raw: return SimpleDF([], [])
        cols = list(raw[0].keys())
        rows = [{k: _try_num(v) for k,v in r.items()} for r in raw]
        return SimpleDF(rows, cols)
    else:
        import openpyxl as _oxl
        wb = _oxl.load_workbook(path, data_only=True)
        ws = wb.active
        data = list(ws.values)
        if not data: return SimpleDF([], [])
        headers = [str(h) if h is not None else f"col_{i}" for i,h in enumerate(data[0])]
        rows = [dict(zip(headers, row)) for row in data[1:]]
        return SimpleDF(rows, headers)

def validate(f1_path, f2_path, prog):
    prog("📥  Membaca File 1 …")
    df1 = load_df(f1_path)
    req = ["product_id","product_name","l1_category_name",
           "pricing_bl_25","qty","selling_price","cost_price"]
    miss = [c for c in req if c not in df1.columns]
    if miss: sys.exit(f"❌  Kolom File 1 tidak lengkap: {miss}")

    # clean product_id
    import re as _re
    df1["product_id"] = [_re.sub(r"\.0$","",str(v).strip()) for v in df1["product_id"]]
    # l2 boleh kosong
    if "l2_category_name" not in df1.columns:
        df1["l2_category_name"] = [""] * len(df1)
    df1["l2_category_name"] = df1.fillna("", col="l2_category_name")["l2_category_name"]

    dup_mask = df1.duplicated(keep=False, cols=["product_id"])
    if df1.any_true(dup_mask):
        all_ids = [r.get("product_id") for r,d in zip(df1._data,dup_mask) if d]
        sys.exit(f"❌  Duplicate product_id di File 1: {list(set(all_ids))[:5]}")

    for col in req:
        if df1.isna_any(col):
            sys.exit(f"❌  Kolom '{col}' di File 1 ada nilai kosong.")

    prog("📥  Membaca File 2 …")
    df2r = load_df(f2_path)
    import re as _re
    col0 = [_re.sub(r"\.0$","",str(v).strip()) if v is not None else "" for v in df2r["product_id"]]
    # rebuild df2r with string col0
    new_rows = []
    for i, r in enumerate(df2r._data):
        new_r = {df2r._cols[0]: col0[i]}
        for c in df2r._cols[1:]:
            new_r[c] = r.get(c)
        new_rows.append(new_r)
    df2r = SimpleDF(new_rows, df2r._cols)
    n = df2r.shape[1]
    if n < 2: sys.exit("❌  File 2 minimal 2 kolom.")
    df2r._cols = ["product_id","baseline"] + [f"var_{i}" for i in range(1,n-1)]
    for r in df2r._data:
        old_keys = list(r.keys())
        new_keys = df2r._cols
        new_r = {nk: r.get(ok) for ok,nk in zip(old_keys, new_keys)}
        r.clear(); r.update(new_r)
    df2 = df2r.dropna(subset=["product_id"]).copy()

    for col in df2._cols[1:]:
        if df2.isna_any(col):
            sys.exit(f"❌  Kolom '{col}' di File 2 ada nilai kosong.")

    ids1 = set(df1["product_id"])
    bad  = [i for i in df2["product_id"] if i not in ids1]
    if bad: sys.exit(f"❌  product_id di File 2 tidak ada di File 1: {bad[:5]}")

    vcols = [c for c in df2.columns if c.startswith("var_")]
    prog(f"   ✅  {len(df2)} SKU akan diubah | {len(vcols)} variant")
    return df1, df2, vcols

# ── SCOPE ─────────────────────────────────────────────────────────────────────
def detect_scope(df1, price_ids):
    sub   = df1.where_isin("product_id", price_ids)
    scope = {}
    for lv, col in [("bl","pricing_bl_25"),("l1","l1_category_name")]:
        vals = sub.unique(col)
        if len(vals) == 1:
            scope[lv] = vals[0]
    return scope

# ── STYLE HELPERS ─────────────────────────────────────────────────────────────
def write_header_cell(ws, r, c, val, bg, font_color=None, wrap=False, size=9, bold=True):
    cell = ws.cell(r, c, val)
    cell.fill      = fx(bg)
    cell.font      = ft(bold=bold, color=font_color or C["font_white"], size=size)
    cell.alignment = al("center","center", wrap=wrap)
    cell.border    = bdr(C["border_dark"])
    return cell

def write_data_cell(ws, r, c, val, bg, font_color=None, fmt=None,
                    halign="right", bold=False, italic=False):
    cell = ws.cell(r, c, val)
    cell.fill      = fx(bg)
    cell.font      = ft(bold=bold, color=font_color or C["font_dark"],
                        italic=italic)
    cell.alignment = al(halign, "center")
    cell.border    = bdr()
    if fmt: cell.number_format = fmt
    return cell

# ── SHEET 1 — RAW DATA ────────────────────────────────────────────────────────
def build_sheet1(ws, df1, price_ids, prog):
    prog("📊  Sheet 1: Raw Data …")
    ws.sheet_view.showGridLines = False

    COLS = ["is_price_adjusted","product_id","product_name",
            "l1_category_name","l2_category_name","pricing_bl_25",
            "qty","selling_price","cost_price","gv","cogs","gp","gp_pct"]
    LABELS = {
        "is_price_adjusted": "Price\nAdjusted?",
        "product_id":        "Product ID",
        "product_name":      "Product Name",
        "l1_category_name":  "L1 Category",
        "l2_category_name":  "L2 Category",
        "pricing_bl_25":     "Business Line",
        "qty":               "Qty Sold",
        "selling_price":     "Selling Price",
        "cost_price":        "COGS / Unit",
        "gv":                "Goods Value",
        "cogs":              "COGS Total",
        "gp":                "Gross Profit",
        "gp_pct":            "GP%",
    }
    WIDTHS = {
        "is_price_adjusted":14,"product_id":14,"product_name":30,
        "l1_category_name":22,"l2_category_name":20,"pricing_bl_25":14,
        "qty":10,"selling_price":14,"cost_price":14,
        "gv":16,"cogs":16,"gp":16,"gp_pct":10,
    }
    ci = {c:i+1 for i,c in enumerate(COLS)}

    # header
    for c, idx in ci.items():
        bg = C["actual_bg"] if c in ["gv","cogs","gp","gp_pct"] else C["navy"]
        fc = C["font_navy"] if c in ["gv","cogs","gp","gp_pct"] else C["font_white"]
        write_header_cell(ws,1,idx,LABELS[c],bg,fc,wrap=True,size=9)
        ws.column_dimensions[get_column_letter(idx)].width = WIDTHS[c]
    ws.row_dimensions[1].height = 32

    # data
    for i, row in enumerate(df1.itertuples(index=False), start=2):
        pid     = str(getattr(row,"product_id"))
        is_adj  = pid in price_ids
        bg_flag = C["flag_true"] if is_adj else C["flag_false"]

        write_data_cell(ws,i,ci["is_price_adjusted"],
                        "TRUE" if is_adj else "FALSE",
                        bg_flag,
                        C["font_green"] if is_adj else C["font_grey"],
                        halign="center", bold=is_adj)

        for c in ["product_id","product_name","l1_category_name",
                  "l2_category_name","pricing_bl_25"]:
            write_data_cell(ws,i,ci[c],getattr(row,c,None),
                            C["white"], C["font_navy"], halign="left")

        write_data_cell(ws,i,ci["qty"],getattr(row,"qty"),C["white"],fmt=NUM0)
        write_data_cell(ws,i,ci["selling_price"],getattr(row,"selling_price"),
                        C["white"],fmt=NUM0)
        write_data_cell(ws,i,ci["cost_price"],getattr(row,"cost_price"),
                        C["white"],fmt=NUM0)

        # Excel formulas
        r       = i
        qty_c   = get_column_letter(ci["qty"])
        sp_c    = get_column_letter(ci["selling_price"])
        cp_c    = get_column_letter(ci["cost_price"])
        gv_c    = get_column_letter(ci["gv"])
        cogs_c  = get_column_letter(ci["cogs"])
        gp_c    = get_column_letter(ci["gp"])

        write_data_cell(ws,i,ci["gv"],   f"={sp_c}{r}*{qty_c}{r}",
                        C["actual_bg"],fmt=NUM0)
        write_data_cell(ws,i,ci["cogs"], f"={cp_c}{r}*{qty_c}{r}",
                        C["actual_bg"],fmt=NUM0)
        write_data_cell(ws,i,ci["gp"],   f"={gv_c}{r}-{cogs_c}{r}",
                        C["actual_bg"],fmt=NUM0)
        write_data_cell(ws,i,ci["gp_pct"],f"=IF({gv_c}{r}=0,0,{gp_c}{r}/{gv_c}{r})",
                        C["actual_bg"],fmt=PCT)
        ws.row_dimensions[i].height = 16

    ws.freeze_panes = "B2"
    return ci, len(df1)+1  # last data row

# ── SHEET 2 — SKU DETAIL ──────────────────────────────────────────────────────
def build_sheet2(ws, df1, df2, vcols, price_ids, prog):
    prog("📊  Sheet 2: SKU Detail …")
    ws.sheet_view.showGridLines = False

    scenarios = ["baseline"] + vcols  # e.g. baseline, var_1, var_2
    dvars     = vcols                  # only variant cols for delta

    # ── column registry ───────────────────────────────────────────────────
    INFO  = ["product_id","product_name","l1_category_name",
             "l2_category_name","pricing_bl_25"]
    ACT_HIDDEN = ["act_price","act_gv","act_cogs","act_gp","act_gp_pct"]
    INPUT = ["qty","cost_price"]
    PRICE = scenarios
    GV    = [f"gv_{s}"     for s in scenarios]
    COGS  = ["cogs_total"]
    GP    = [f"gp_{s}"     for s in scenarios]
    GPPCT = [f"gp_pct_{s}" for s in scenarios]
    DHARGA  = [f"dh_{v}"   for v in dvars]
    DHARGAP = [f"dhp_{v}"  for v in dvars]
    DGV     = [f"dgv_{v}"  for v in dvars]
    DGVP    = [f"dgvp_{v}" for v in dvars]
    DGP     = [f"dgp_{v}"  for v in dvars]
    DGPP    = [f"dgpp_{v}" for v in dvars]
    DGPPCT  = [f"dgp_pct_{v}" for v in dvars]

    ALL = (ACT_HIDDEN + INFO + INPUT + PRICE + GV + COGS +
           GP + GPPCT + DHARGA + DHARGAP + DGV + DGVP + DGP + DGPP + DGPPCT)
    ci = {c:i+1 for i,c in enumerate(ALL)}

    def cl(name): return get_column_letter(ci[name])

    slabel = lambda s: "Baseline" if s=="baseline" else s.replace("var_","Var ")

    # ── group header row 1 ────────────────────────────────────────────────
    def span(label, start, end, bg):
        sc, ec = ci[start], ci[end]
        write_header_cell(ws,1,sc,label,bg,size=9,bold=True)
        if ec>sc:
            ws.merge_cells(start_row=1,start_column=sc,
                           end_row=1,  end_column=ec)
        for x in range(sc+1,ec+1):
            ws.cell(1,x).fill   = fx(bg)
            ws.cell(1,x).border = bdr(C["border_dark"])

    for c in ACT_HIDDEN:
        ws.cell(1,ci[c]).fill   = fx(C["actual_bg"])
        ws.cell(1,ci[c]).border = bdr(C["border_dark"])

    span("Info SKU",      INFO[0],    INFO[-1],    C["navy"])
    span("Input",         INPUT[0],   INPUT[-1],   C["navy"])
    span("Harga",         PRICE[0],   PRICE[-1],   C["navy_mid"])
    span("Goods Value",   GV[0],      GV[-1],      C["navy_mid"])
    span("COGS",          COGS[0],    COGS[-1],    C["navy_mid"])
    span("Gross Profit",  GP[0],      GP[-1],      C["navy_mid"])
    span("GP%",           GPPCT[0],   GPPCT[-1],   C["navy_mid"])
    if dvars:
        span("Δ Harga",   DHARGA[0],  DHARGA[-1],  "5D4E8E")
        span("Δ% Harga",  DHARGAP[0], DHARGAP[-1], "5D4E8E")
        span("Δ GV",      DGV[0],     DGV[-1],     "5D4E8E")
        span("Δ% GV",     DGVP[0],    DGVP[-1],    "5D4E8E")
        span("Δ GP",      DGP[0],     DGP[-1],     "5D4E8E")
        span("Δ% GP",     DGPP[0],    DGPP[-1],    "5D4E8E")
        span("Δ GP%",     DGPPCT[0],  DGPPCT[-1],  "5D4E8E")
    ws.row_dimensions[1].height = 22

    # ── sub-header row 2 ──────────────────────────────────────────────────
    sub_labels = {
        "act_price":"Act Price","act_gv":"Act GV","act_cogs":"Act COGS",
        "act_gp":"Act GP","act_gp_pct":"Act GP%",
        "product_id":"Product ID","product_name":"Product Name",
        "l1_category_name":"L1 Category","l2_category_name":"L2 Category",
        "pricing_bl_25":"Business Line",
        "qty":"Qty Sold","cost_price":"COGS/Unit","cogs_total":"COGS Total",
    }
    for s in scenarios:
        sub_labels[s]            = f"Harga {slabel(s)}"
        sub_labels[f"gv_{s}"]    = f"GV {slabel(s)}"
        sub_labels[f"gp_{s}"]    = f"GP {slabel(s)}"
        sub_labels[f"gp_pct_{s}"]= f"GP% {slabel(s)}"
    for v in dvars:
        lv = v.replace("var_","Var ")
        sub_labels[f"dh_{v}"]      = f"Δ Harga {lv}"
        sub_labels[f"dhp_{v}"]     = f"Δ% Harga {lv}"
        sub_labels[f"dgv_{v}"]     = f"Δ GV {lv}"
        sub_labels[f"dgvp_{v}"]    = f"Δ% GV {lv}"
        sub_labels[f"dgp_{v}"]     = f"Δ GP {lv}"
        sub_labels[f"dgpp_{v}"]    = f"Δ% GP {lv}"
        sub_labels[f"dgp_pct_{v}"] = f"Δ GP% {lv}"

    for cname, cidx in ci.items():
        is_act = cname in ACT_HIDDEN
        bg = C["actual_bg"] if is_act else C["header_sub"]
        fc = C["font_grey"] if is_act else C["font_navy"]
        write_header_cell(ws,2,cidx,sub_labels.get(cname,cname),
                          bg,fc,wrap=True,size=8,bold=True)
    ws.row_dimensions[2].height = 32

    # ── data rows ─────────────────────────────────────────────────────────
    # merge df1 info for price_ids
    df2m = df2.merge(
        df1,
        on="product_id", how="left"
    )

    DATA_START = 3
    for i, row in enumerate(df2m.itertuples(index=False), start=DATA_START):
        r = i
        # hidden actual cols
        for c in ACT_HIDDEN:
            ws.cell(r,ci[c]).fill   = fx(C["actual_bg"])
            ws.cell(r,ci[c]).border = bdr()

        # info (left-aligned, navy text)
        for c in INFO:
            write_data_cell(ws,r,ci[c],getattr(row,c,None),
                            C["white"],C["font_navy"],halign="left")

        # input (hardcoded, blue text = editable signal)
        write_data_cell(ws,r,ci["qty"],      getattr(row,"qty"),
                        C["input_bg"],C["font_input"],fmt=NUM0)
        write_data_cell(ws,r,ci["cost_price"],getattr(row,"cost_price"),
                        C["input_bg"],C["font_input"],fmt=NUM0)

        # harga — hardcoded from File 2 (editable)
        write_data_cell(ws,r,ci["baseline"],getattr(row,"baseline"),
                        C["baseline_bg"],C["font_input"],fmt=NUM0,bold=False)
        for v in dvars:
            write_data_cell(ws,r,ci[v],getattr(row,v,None),
                            C["var_bg"],C["font_input"],fmt=NUM0)

        # column letter shortcuts
        qty_c  = cl("qty")
        cp_c   = cl("cost_price")
        bl_c   = cl("baseline")
        ct_c   = cl("cogs_total")

        # actual hidden formulas
        sp_val = getattr(row,"selling_price")
        ws.cell(r,ci["act_price"]).value          = sp_val
        ws.cell(r,ci["act_price"]).number_format  = NUM0
        ws.cell(r,ci["act_gv"]).value             = f"={cl('act_price')}{r}*{qty_c}{r}"
        ws.cell(r,ci["act_gv"]).number_format     = NUM0
        ws.cell(r,ci["act_cogs"]).value           = f"={cp_c}{r}*{qty_c}{r}"
        ws.cell(r,ci["act_cogs"]).number_format   = NUM0
        ws.cell(r,ci["act_gp"]).value             = f"={cl('act_gv')}{r}-{cl('act_cogs')}{r}"
        ws.cell(r,ci["act_gp"]).number_format     = NUM0
        ws.cell(r,ci["act_gp_pct"]).value         = f"=IF({cl('act_gv')}{r}=0,0,{cl('act_gp')}{r}/{cl('act_gv')}{r})"
        ws.cell(r,ci["act_gp_pct"]).number_format = PCT
        for c in ACT_HIDDEN:
            ws.cell(r,ci[c]).fill   = fx(C["actual_bg"])
            ws.cell(r,ci[c]).border = bdr()
            ws.cell(r,ci[c]).font   = ft(color=C["font_grey"],italic=True)

        # GV formulas
        for s in scenarios:
            harga_c = cl(s)
            gv_c    = cl(f"gv_{s}")
            ws.cell(r,ci[f"gv_{s}"]).value          = f"={harga_c}{r}*{qty_c}{r}"
            ws.cell(r,ci[f"gv_{s}"]).number_format  = NUM0
            ws.cell(r,ci[f"gv_{s}"]).fill           = fx(C["baseline_bg"] if s=="baseline" else C["var_bg"])
            ws.cell(r,ci[f"gv_{s}"]).border         = bdr()
            ws.cell(r,ci[f"gv_{s}"]).alignment      = al("right","center")

        # COGS Total
        ws.cell(r,ci["cogs_total"]).value          = f"={cp_c}{r}*{qty_c}{r}"
        ws.cell(r,ci["cogs_total"]).number_format  = NUM0
        ws.cell(r,ci["cogs_total"]).fill           = fx(C["white"])
        ws.cell(r,ci["cogs_total"]).border         = bdr()
        ws.cell(r,ci["cogs_total"]).alignment      = al("right","center")

        # GP formulas
        for s in scenarios:
            gv_c = cl(f"gv_{s}")
            ws.cell(r,ci[f"gp_{s}"]).value          = f"={gv_c}{r}-{ct_c}{r}"
            ws.cell(r,ci[f"gp_{s}"]).number_format  = NUM0
            ws.cell(r,ci[f"gp_{s}"]).fill           = fx(C["baseline_bg"] if s=="baseline" else C["var_bg"])
            ws.cell(r,ci[f"gp_{s}"]).border         = bdr()
            ws.cell(r,ci[f"gp_{s}"]).alignment      = al("right","center")

        # GP% formulas
        for s in scenarios:
            gv_c = cl(f"gv_{s}")
            gp_c = cl(f"gp_{s}")
            ws.cell(r,ci[f"gp_pct_{s}"]).value          = f"=IF({gv_c}{r}=0,0,{gp_c}{r}/{gv_c}{r})"
            ws.cell(r,ci[f"gp_pct_{s}"]).number_format  = PCT
            ws.cell(r,ci[f"gp_pct_{s}"]).fill           = fx(C["baseline_bg"] if s=="baseline" else C["var_bg"])
            ws.cell(r,ci[f"gp_pct_{s}"]).border         = bdr()
            ws.cell(r,ci[f"gp_pct_{s}"]).alignment      = al("right","center")

        # Delta formulas (vs Baseline)
        for v in dvars:
            hv_c  = cl(v)
            gvv_c = cl(f"gv_{v}")
            gpv_c = cl(f"gp_{v}")
            gppv_c= cl(f"gp_pct_{v}")
            gvbl  = cl("gv_baseline")
            gpbl  = cl("gp_baseline")
            gppbl = cl("gp_pct_baseline")

            def dw(cname, formula, fmt, bg=C["delta_bg"]):
                ws.cell(r,ci[cname]).value          = formula
                ws.cell(r,ci[cname]).number_format  = fmt
                ws.cell(r,ci[cname]).fill           = fx(bg)
                ws.cell(r,ci[cname]).border         = bdr()
                ws.cell(r,ci[cname]).alignment      = al("right","center")

            dw(f"dh_{v}",      f"={hv_c}{r}-{bl_c}{r}",              DNUM)
            dw(f"dhp_{v}",     f"=IF({bl_c}{r}=0,0,({hv_c}{r}-{bl_c}{r})/{bl_c}{r})", DPCT)
            dw(f"dgv_{v}",     f"={gvv_c}{r}-{gvbl}{r}",             DNUM)
            dw(f"dgvp_{v}",    f"=IF({gvbl}{r}=0,0,({gvv_c}{r}-{gvbl}{r})/{gvbl}{r})",DPCT)
            dw(f"dgp_{v}",     f"={gpv_c}{r}-{gpbl}{r}",             DNUM)
            dw(f"dgpp_{v}",    f"=IF({gpbl}{r}=0,0,({gpv_c}{r}-{gpbl}{r})/{gpbl}{r})",DPCT)
            dw(f"dgp_pct_{v}", f"={gppv_c}{r}-{gppbl}{r}",           DPP)

        ws.row_dimensions[r].height = 16

    # hide actual cols
    for c in ACT_HIDDEN:
        ws.column_dimensions[get_column_letter(ci[c])].hidden = True

    # col widths
    w_map = {
        "product_id":14,"product_name":28,"l1_category_name":22,
        "l2_category_name":20,"pricing_bl_25":14,"qty":10,"cost_price":12,
        "cogs_total":14,
    }
    for cname in ALL:
        w = w_map.get(cname,14)
        ws.column_dimensions[get_column_letter(ci[cname])].width = w

    ws.freeze_panes = f"{get_column_letter(ci['product_id'])}3"
    return ci, DATA_START, DATA_START + len(df2m) - 1  # first_row, last_row

# ── SHEET 3 — SUMMARY IMPACT ─────────────────────────────────────────────────
def build_sheet3(ws, df1, df2, vcols, price_ids, scope,
                 s1_ci, s1_last,
                 s2_ci, s2_first, s2_last, prog):
    prog("📊  Sheet 3: Summary Impact …")
    ws.sheet_view.showGridLines = False

    dvars     = vcols
    scenarios = ["baseline"] + dvars
    slabel    = lambda s: "Baseline" if s=="baseline" else s.replace("var_","Var ")
    n_vars    = len(dvars)

    # column layout: Metrik | Satuan | Actual* | Baseline | Var1..N | ΔVar1..N | Δ%Var1..N
    COL_METRIC    = 1
    COL_SATUAN    = 2
    COL_ACTUAL    = 3
    COL_BASELINE  = 4
    COL_VAR_START = 5
    COL_D_START   = COL_VAR_START + n_vars
    COL_DP_START  = COL_D_START   + n_vars
    TOTAL_COLS    = COL_DP_START  + n_vars - 1

    # helper: get column letter in Sheet2 by key
    def s2cl(key): return get_column_letter(s2_ci[key])
    def s1cl(key): return get_column_letter(s1_ci[key])

    # Sheet names
    S1 = "'Raw Data'"
    S2 = "'SKU Detail'"

    # SUMIF ranges for Sheet 1 (non-price-change SKUs)
    # flag col in Sheet 1
    flag_col_s1 = s1cl("is_price_adjusted")
    s1_range    = f"{S1}!{flag_col_s1}2:{flag_col_s1}{s1_last}"

    def s1_sum_col(metric_col, extra_filter_col=None, extra_filter_val=None):
        """SUMIF/SUMPRODUCT on Sheet1 for FALSE rows, optionally filtered by level"""
        val_col = f"{S1}!{s1cl(metric_col)}2:{s1cl(metric_col)}{s1_last}"
        if extra_filter_col and extra_filter_val:
            lvl_col = f"{S1}!{s1cl(extra_filter_col)}2:{s1cl(extra_filter_col)}{s1_last}"
            return (f'SUMPRODUCT(({s1_range}=FALSE)*'
                    f'({lvl_col}="{extra_filter_val}")*'
                    f'{val_col})')
        return f'SUMIF({s1_range},FALSE,{val_col})'

    # SUMIF on Sheet 2 (price-change SKUs)
    def s2_sum_col(metric_col, filter_col=None, filter_val=None):
        """SUM/SUMIF on Sheet2"""
        val_col = f"{S2}!{s2cl(metric_col)}{s2_first}:{s2cl(metric_col)}{s2_last}"
        if filter_col and filter_val:
            lvl_col = f"{S2}!{s2cl(filter_col)}{s2_first}:{s2cl(filter_col)}{s2_last}"
            return f'SUMIF({lvl_col},"{filter_val}",{val_col})'
        return f'SUM({val_col})'

    def combined_gv(s, fc=None, fv=None):
        s2_gv = s2_sum_col(f"gv_{s}", fc, fv)
        s1_gv = s1_sum_col("gv", fc, fv)
        return f"={s2_gv}+{s1_gv}"

    def combined_cogs(fc=None, fv=None):
        s2_c = s2_sum_col("cogs_total", fc, fv)
        s1_c = s1_sum_col("cogs", fc, fv)
        return f"={s2_c}+{s1_c}"

    def combined_qty(fc=None, fv=None):
        s2_q = s2_sum_col("qty", fc, fv)
        s1_q = s1_sum_col("qty", fc, fv)
        return f"={s2_q}+{s1_q}"

    def combined_actual_gv(fc=None, fv=None):
        s2_a = s2_sum_col("act_gv", fc, fv)
        s1_a = s1_sum_col("gv", fc, fv)
        return f"={s2_a}+{s1_a}"

    def combined_actual_gp(fc=None, fv=None):
        s2_a = s2_sum_col("act_gp", fc, fv)
        s1_a = s1_sum_col("gp", fc, fv)
        return f"={s2_a}+{s1_a}"

    # ── render tables ─────────────────────────────────────────────────────
    cur_row = 1

    def write_table(title, fc=None, fv=None):
        nonlocal cur_row

        # title row
        ws.merge_cells(start_row=cur_row,start_column=1,
                       end_row=cur_row,  end_column=TOTAL_COLS)
        tc = ws.cell(cur_row,1,title)
        tc.fill      = fx(C["s3_title"])
        tc.font      = ft(bold=True,color=C["font_white"],size=10)
        tc.alignment = al("left","center")
        tc.border    = bdr(C["border_dark"])
        ws.row_dimensions[cur_row].height = 22
        cur_row += 1

        # sub-header
        headers = ["Metrik","Satuan","Actual*","Baseline"]
        for v in dvars: headers.append(slabel(v))
        for v in dvars: headers.append(f"Δ {slabel(v)}")
        for v in dvars: headers.append(f"Δ% {slabel(v)}")

        for ci2, h in enumerate(headers,1):
            is_act = ci2 == COL_ACTUAL
            bg = C["actual_bg"] if is_act else C["header_sub"]
            fc2= C["font_grey"] if is_act else C["font_navy"]
            write_header_cell(ws,cur_row,ci2,h,bg,fc2,wrap=True,size=8)
        ws.row_dimensions[cur_row].height = 28
        cur_row += 1

        # METRICS: (label, satuan, gv_key or special, fmt, is_gp_pct)
        METRICS = [
            ("Goods Value", "Rp",   "gv",     NUM0,  False),
            ("Qty",         "Unit", "qty",    NUM0,  False),
            ("COGS",        "Rp",   "cogs",   NUM0,  False),
            ("GP",          "Rp",   "gp",     NUM0,  False),
            ("GP%",         "%",    "gp_pct", PCT,   True),
        ]

        for mlabel, satuan, key, fmt, is_pct in METRICS:
            r = cur_row
            write_data_cell(ws,r,COL_METRIC,mlabel,C["s3_metric"],
                            C["font_navy"],halign="left",bold=True)
            write_data_cell(ws,r,COL_SATUAN,satuan,C["s3_metric"],
                            C["font_grey"],halign="center")

            # Actual (hidden)
            if key == "gv":
                act_f = combined_actual_gv(fc,fv)
            elif key == "gp":
                act_f = combined_actual_gp(fc,fv)
            elif key == "qty":
                act_f = combined_qty(fc,fv)
            elif key == "cogs":
                act_f = combined_cogs(fc,fv)
            else:  # gp_pct
                act_gv_f = combined_actual_gv(fc,fv).lstrip("=")
                act_gp_f = combined_actual_gp(fc,fv).lstrip("=")
                act_f = f"=IF(({act_gv_f})=0,0,({act_gp_f})/({act_gv_f}))"

            ac = ws.cell(r,COL_ACTUAL,act_f)
            ac.fill=fx(C["actual_bg"]); ac.border=bdr()
            ac.number_format=fmt; ac.alignment=al("right","center")
            ac.font=ft(color=C["font_grey"],italic=True)

            # Baseline
            if key == "gv":
                bl_f = combined_gv("baseline",fc,fv)
            elif key == "qty":
                bl_f = combined_qty(fc,fv)
            elif key == "cogs":
                bl_f = combined_cogs(fc,fv)
            elif key == "gp":
                gv_f  = combined_gv("baseline",fc,fv).lstrip("=")
                cog_f = combined_cogs(fc,fv).lstrip("=")
                bl_f  = f"=({gv_f})-({cog_f})"
            else:  # gp_pct
                gv_f  = combined_gv("baseline",fc,fv).lstrip("=")
                cog_f = combined_cogs(fc,fv).lstrip("=")
                bl_f  = f"=IF(({gv_f})=0,0,(({gv_f})-({cog_f}))/({gv_f}))"

            bc = ws.cell(r,COL_BASELINE,bl_f)
            bc.fill=fx(C["s3_baseline"]); bc.border=bdr()
            bc.number_format=fmt; bc.alignment=al("right","center")

            # Variants
            for vi, v in enumerate(dvars):
                col = COL_VAR_START + vi
                if key == "gv":
                    vf = combined_gv(v,fc,fv)
                elif key == "qty":
                    vf = combined_qty(fc,fv)
                elif key == "cogs":
                    vf = combined_cogs(fc,fv)
                elif key == "gp":
                    gv_f  = combined_gv(v,fc,fv).lstrip("=")
                    cog_f = combined_cogs(fc,fv).lstrip("=")
                    vf    = f"=({gv_f})-({cog_f})"
                else:  # gp_pct
                    gv_f  = combined_gv(v,fc,fv).lstrip("=")
                    cog_f = combined_cogs(fc,fv).lstrip("=")
                    vf    = f"=IF(({gv_f})=0,0,(({gv_f})-({cog_f}))/({gv_f}))"

                vc = ws.cell(r,col,vf)
                vc.fill=fx(C["s3_var"]); vc.border=bdr()
                vc.number_format=fmt; vc.alignment=al("right","center")

            # Deltas (vs Baseline)
            bl_col_l = get_column_letter(COL_BASELINE)
            for vi, v in enumerate(dvars):
                var_col_l = get_column_letter(COL_VAR_START+vi)
                d_col     = COL_D_START  + vi
                dp_col    = COL_DP_START + vi

                if is_pct:
                    df_formula = f"={var_col_l}{r}-{bl_col_l}{r}"
                    dp_formula = '="—"'
                    df_fmt     = DPP
                    dp_fmt     = "@"
                else:
                    df_formula = f"={var_col_l}{r}-{bl_col_l}{r}"
                    dp_formula = f"=IF({bl_col_l}{r}=0,0,({var_col_l}{r}-{bl_col_l}{r})/{bl_col_l}{r})"
                    df_fmt     = DNUM
                    dp_fmt     = DPCT

                dc = ws.cell(r,d_col,df_formula)
                dc.fill=fx(C["s3_delta"]); dc.border=bdr()
                dc.number_format=df_fmt; dc.alignment=al("right","center")

                dpc = ws.cell(r,dp_col,dp_formula)
                dpc.fill=fx(C["s3_delta"]); dpc.border=bdr()
                dpc.number_format=dp_fmt; dpc.alignment=al("right","center")

            ws.row_dimensions[r].height = 18
            cur_row += 1

        cur_row += 2  # gap

    # ── define tables ─────────────────────────────────────────────────────
    n_price = len(price_ids)
    write_table("OVERALL")

    if n_price == 1:
        row0 = df1.where_isin("product_id", price_ids)._data[0]
        write_table(f"BUSINESS LINE: {row0['pricing_bl_25']}",
                    "pricing_bl_25", row0["pricing_bl_25"])
        write_table(f"L1 CATEGORY: {row0['l1_category_name']}",
                    "l1_category_name", row0["l1_category_name"])
    else:
        sub = df1.where_isin("product_id", price_ids)
        for col, prefix in [("pricing_bl_25","BUSINESS LINE"),
                             ("l1_category_name","L1 CATEGORY")]:
            vals = sub.unique(col)
            if len(vals) == 1:
                write_table(f"{prefix}: {vals[0]}", col, vals[0])

    write_table("ALL SKU PRICE ADJUST")

    # hide actual col
    ws.column_dimensions[get_column_letter(COL_ACTUAL)].hidden = True

    # col widths
    ws.column_dimensions[get_column_letter(COL_METRIC)].width   = 14
    ws.column_dimensions[get_column_letter(COL_SATUAN)].width   = 7
    ws.column_dimensions[get_column_letter(COL_ACTUAL)].width   = 14
    ws.column_dimensions[get_column_letter(COL_BASELINE)].width = 16
    for i in range(n_vars):
        ws.column_dimensions[get_column_letter(COL_VAR_START+i)].width   = 16
        ws.column_dimensions[get_column_letter(COL_D_START+i)].width     = 14
        ws.column_dimensions[get_column_letter(COL_DP_START+i)].width    = 12

# ── MAIN ──────────────────────────────────────────────────────────────────────
def main():
    def prog(msg): print(f"  {msg}")

    f1_path, f2_path = pick_files()
    df1, df2, vcols  = validate(f1_path, f2_path, prog)
    price_ids        = set(df2["product_id"])
    scope            = detect_scope(df1, price_ids)

    prog("📝  Membuat workbook …")
    wb   = openpyxl.Workbook()
    ws1  = wb.active;              ws1.title = "Raw Data"
    ws2  = wb.create_sheet("SKU Detail")
    ws3  = wb.create_sheet("Summary Impact")

    s1_ci, s1_last               = build_sheet1(ws1, df1, price_ids, prog)
    s2_ci, s2_first, s2_last     = build_sheet2(ws2, df1, df2, vcols, price_ids, prog)
    build_sheet3(ws3, df1, df2, vcols, price_ids, scope,
                 s1_ci, s1_last,
                 s2_ci, s2_first, s2_last, prog)

    out = os.path.join(os.path.dirname(f1_path), "pricing_impact_output.xlsx")
    prog(f"💾  Menyimpan ke {out} …")
    wb.save(out)
    prog(f"✅  Selesai! → {out}")
    print(f"\n{'='*55}\n  OUTPUT: {out}\n{'='*55}\n")




def validate_bytes(f1_bytes, f1_name, f2_bytes, f2_name, prog):
    """Pyodide-compatible version of validate() that accepts file bytes"""
    import io, csv as _csv, re as _re

    prog("Reading File 1...")
    ext1 = f1_name.rsplit(".", 1)[-1].lower() if "." in f1_name else "csv"
    if ext1 == "csv":
        text = f1_bytes.decode("utf-8-sig")
        reader = _csv.DictReader(io.StringIO(text))
        raw = list(reader)
        if not raw: raise ValueError("File 1 kosong")
        cols = list(raw[0].keys())
        rows = [{k: _try_num(v) for k,v in r.items()} for r in raw]
        df1 = SimpleDF(rows, cols)
    else:
        import openpyxl as _oxl
        wb_tmp = _oxl.load_workbook(io.BytesIO(f1_bytes), read_only=True, data_only=True)
        ws_tmp = wb_tmp.active
        rows_raw = list(ws_tmp.values)
        wb_tmp.close()
        if not rows_raw: raise ValueError("File 1 kosong")
        headers = [str(h) if h is not None else "" for h in rows_raw[0]]
        data_rows = [{headers[i]: _try_num(v) for i,v in enumerate(r)}
                     for r in rows_raw[1:] if any(v is not None for v in r)]
        df1 = SimpleDF(data_rows, headers)

    req = ["product_id","product_name","l1_category_name","pricing_bl_25","qty","selling_price","cost_price"]
    miss = [c for c in req if c not in df1.columns]
    if miss: raise ValueError(f"Kolom File 1 tidak lengkap: {miss}")

    for i, r in enumerate(df1._data):
        df1._data[i]["product_id"] = _re.sub(r"\.0$", "", str(r.get("product_id","")).strip())

    if "l2_category_name" not in df1.columns:
        df1._cols.append("l2_category_name")
        for r in df1._data: r["l2_category_name"] = ""

    prog("Reading File 2...")
    ext2 = f2_name.rsplit(".", 1)[-1].lower() if "." in f2_name else "csv"
    if ext2 == "csv":
        text2 = f2_bytes.decode("utf-8-sig")
        reader2 = _csv.DictReader(io.StringIO(text2))
        raw2 = list(reader2)
        if not raw2: raise ValueError("File 2 kosong")
        cols2 = list(raw2[0].keys())
        rows2 = [{k: _try_num(v) for k,v in r.items()} for r in raw2]
        df2r = SimpleDF(rows2, cols2)
    else:
        import openpyxl as _oxl
        wb_tmp = _oxl.load_workbook(io.BytesIO(f2_bytes), read_only=True, data_only=True)
        ws_tmp = wb_tmp.active
        rows_raw2 = list(ws_tmp.values)
        wb_tmp.close()
        if not rows_raw2: raise ValueError("File 2 kosong")
        headers2 = [str(h) if h is not None else "" for h in rows_raw2[0]]
        data_rows2 = [{headers2[i]: _try_num(v) for i,v in enumerate(r)}
                      for r in rows_raw2[1:] if any(v is not None for v in r)]
        df2r = SimpleDF(data_rows2, headers2)

    for i, r in enumerate(df2r._data):
        df2r._data[i]["product_id"] = _re.sub(r"\.0$", "", str(r.get("product_id","")).strip())

    vcols = [c for c in df2r.columns if c != "product_id"]
    if not vcols: raise ValueError("File 2 harus punya minimal kolom baseline")
    if "baseline" not in vcols: raise ValueError("File 2 harus punya kolom baseline")

    ids1 = set(r["product_id"] for r in df1._data)
    df2r._data = [r for r in df2r._data if r["product_id"] in ids1]

    return df1, df2r, vcols


def run_sim(file1_bytes, filename1, file2_bytes, filename2):
    import io, openpyxl

    def prog(msg): print(f"  {msg}")

    df1, df2, vcols = validate_bytes(file1_bytes, filename1, file2_bytes, filename2, prog)
    price_ids = set(r["product_id"] for r in df2._data)
    scope = detect_scope(df1, price_ids)

    prog("Building workbook...")
    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = "Raw Data"
    ws2 = wb.create_sheet("SKU Detail")
    ws3 = wb.create_sheet("Summary Impact")

    s1_ci, s1_last               = build_sheet1(ws1, df1, price_ids, prog)
    s2_ci, s2_first, s2_last     = build_sheet2(ws2, df1, df2, vcols, price_ids, prog)
    build_sheet3(ws3, df1, df2, vcols, price_ids, scope,
                 s1_ci, s1_last,
                 s2_ci, s2_first, s2_last, prog)

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)
    return out_buf.getvalue(), "pricing_impact_output.xlsx"
